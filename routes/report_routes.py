import os
import io
import json

from flask import Blueprint, abort, current_app, jsonify, redirect, render_template, request, send_file, url_for

from database.models import MedicalEntity, MedicalReport, Patient, Reminder, db, safe_json_list
from routes.auth_routes import get_current_user, log_audit, login_required, require_role
from services.doctor_questions import generate_doctor_questions
from services.lab_trends import build_lab_trend
from services.medical_safety import DISCLAIMER
from services.evaluation import evaluate_lab_extraction, evaluate_readability, load_labeled_samples
from services.privacy import delete_report_file
from services.pdf_export import build_simplified_report_pdf
from services.medicine_lookup import lookup_medicine
from services.translation import (
	TranslationNotConfiguredError,
	TranslationProviderError,
	translate_report_output,
)

report_bp = Blueprint("report", __name__)


@report_bp.route("/reports", methods=["GET"])
@login_required
def list_reports():
	reports = MedicalReport.query.filter_by(owner_id=get_current_user().id).order_by(MedicalReport.created_at.desc()).all()
	return render_template("reports.html", reports=reports)


@report_bp.route("/api/reports", methods=["GET"])
@login_required
def api_list_reports():
	reports = MedicalReport.query.filter_by(owner_id=get_current_user().id).order_by(MedicalReport.created_at.desc()).all()
	return jsonify([report.to_dict() for report in reports])


@report_bp.route("/api/reports/search", methods=["GET"])
@login_required
def search_reports():
	user = get_current_user()
	query = request.args.get("q", "").strip().lower()
	condition = request.args.get("condition", "").strip().lower()
	reports = MedicalReport.query.filter_by(owner_id=user.id).order_by(MedicalReport.created_at.desc()).all()
	results = []
	for report in reports:
		search_text = " ".join([
			report.title or "",
			report.filename or "",
			report.extracted_text or "",
			report.patient.full_name if report.patient else "",
		]).lower()
		conditions = " ".join((entity.entity_text or "") for entity in report.entities).lower()
		if query and query not in search_text and query not in conditions:
			continue
		if condition and condition not in conditions:
			continue
		results.append(report.to_dict())
	return jsonify(results)


@report_bp.route("/api/dashboard/stats", methods=["GET"])
@login_required
def dashboard_stats():
	from datetime import datetime, timezone
	user = get_current_user()
	month_start = datetime.now(timezone.utc).replace(day=1, hour=0, minute=0, second=0, microsecond=0)
	reports = MedicalReport.query.filter_by(owner_id=user.id).all()

	def is_from_current_month(report):
		"""Compare SQLite's naive timestamps safely with the UTC month boundary."""
		created_at = report.created_at
		if not created_at:
			return False
		if created_at.tzinfo is None:
			created_at = created_at.replace(tzinfo=timezone.utc)
		return created_at >= month_start

	condition_counts = {}
	for report in reports:
		for entity in report.entities:
			condition_counts[entity.entity_text] = condition_counts.get(entity.entity_text, 0) + 1
	reminders = Reminder.query.filter_by(owner_id=user.id, status="pending").order_by(Reminder.created_at.desc()).limit(6).all()
	return jsonify({
		"patient_count": Patient.query.filter_by(owner_id=user.id).count(),
		"report_count": len(reports),
		"reports_this_month": sum(1 for report in reports if is_from_current_month(report)),
		"most_common_conditions": sorted(condition_counts.items(), key=lambda item: item[1], reverse=True)[:5],
		"reminders": [r.to_dict() for r in reminders],
		"pending_reminders_count": Reminder.query.filter_by(owner_id=user.id, status="pending").count(),
	})


@report_bp.route("/api/reminders", methods=["GET", "POST"])
@login_required
def handle_reminders():
	user = get_current_user()
	if request.method == "POST":
		payload = request.get_json(silent=True) or {}
		title = str(payload.get("title", "")).strip()
		if not title:
			return jsonify({"error": "Reminder title is required."}), 400
		reminder = Reminder(
			owner_id=user.id,
			patient_id=payload.get("patient_id"),
			report_id=payload.get("report_id"),
			title=title,
			due_text=str(payload.get("due_text", "")).strip() or "Scheduled",
			tag=str(payload.get("tag", "info")).strip() or "info",
			notes=str(payload.get("notes", "")).strip(),
			status="pending",
		)
		db.session.add(reminder)
		db.session.commit()
		return jsonify(reminder.to_dict()), 201

	reminders = Reminder.query.filter_by(owner_id=user.id).order_by(
		Reminder.status.asc(), Reminder.created_at.desc()
	).all()
	return jsonify([r.to_dict() for r in reminders])


@report_bp.route("/api/reminders/<int:reminder_id>", methods=["PATCH", "DELETE"])
@login_required
def modify_reminder(reminder_id):
	user = get_current_user()
	reminder = Reminder.query.filter_by(id=reminder_id, owner_id=user.id).first_or_404()
	if request.method == "DELETE":
		db.session.delete(reminder)
		db.session.commit()
		return jsonify({"deleted": True})

	payload = request.get_json(silent=True) or {}
	if "status" in payload:
		reminder.status = "completed" if payload["status"] == "completed" else "pending"
	if "title" in payload and payload["title"].strip():
		reminder.title = payload["title"].strip()
	if "due_text" in payload:
		reminder.due_text = payload["due_text"].strip()
	if "notes" in payload:
		reminder.notes = payload["notes"].strip()
	db.session.commit()
	return jsonify(reminder.to_dict())



@report_bp.route("/api/patients/<int:patient_id>/timeline", methods=["GET"])
@login_required
def patient_timeline(patient_id):
	patient = Patient.query.filter_by(id=patient_id, owner_id=get_current_user().id).first_or_404()
	return jsonify([
		{
			"id": report.id,
			"title": report.title or report.filename or "Medical report",
			"created_at": report.created_at.isoformat() if report.created_at else None,
			"entities": [entity.to_dict() for entity in report.entities],
			"lab_results": report.to_dict()["lab_results"],
			"questions_to_ask_doctor": generate_doctor_questions(report.to_dict()["lab_results"]),
		}
		for report in sorted(patient.reports, key=lambda item: item.created_at or 0, reverse=True)
	])


@report_bp.route("/api/patients/<int:patient_id>/lab-trends", methods=["GET"])
@login_required
def patient_lab_trend(patient_id):
	patient = Patient.query.filter_by(id=patient_id, owner_id=get_current_user().id).first_or_404()
	test_name = request.args.get("test", "").strip()
	if not test_name:
		return jsonify({"error": "A laboratory test name is required."}), 400

	reports = MedicalReport.query.filter_by(
		owner_id=get_current_user().id,
		patient_id=patient.id,
	).order_by(MedicalReport.created_at.asc()).all()
	return jsonify(build_lab_trend([report.to_dict() for report in reports], test_name))


@report_bp.route("/api/reports/<int:report_id>/evaluation", methods=["GET"])
@login_required
def evaluate_report(report_id):
	report = MedicalReport.query.filter_by(id=report_id, owner_id=get_current_user().id).first_or_404()
	return jsonify({
		"readability": {
			"original": evaluate_readability(report.extracted_text or ""),
			"simplified": evaluate_readability(report.simplified_text or ""),
		},
		"lab_extraction": evaluate_lab_extraction(load_labeled_samples()),
	})


@report_bp.route("/api/reports/<int:report_id>/medicine-lookup", methods=["POST"])
@login_required
def lookup_report_medicine(report_id):
	report = MedicalReport.query.filter_by(
		id=report_id,
		owner_id=get_current_user().id,
	).first_or_404()
	payload = request.get_json(silent=True) or {}
	entity_id = payload.get("entity_id")
	name = str(payload.get("name", "")).strip()
	if not isinstance(entity_id, int) or not name or len(name) > 255:
		return jsonify({"error": "A medicine entity and a name of 1 to 255 characters are required."}), 400

	entity = MedicalEntity.query.filter_by(
		id=entity_id,
		report_id=report.id,
	).first_or_404()
	if str(entity.entity_label or "").upper() not in {"MEDICINE", "MEDICATION"}:
		return jsonify({"error": "The selected report entity is not a medicine."}), 400

	other_medicines = [
		item.entity_text
		for item in report.entities
		if item.id != entity.id
		and str(item.entity_label or "").upper() in {"MEDICINE", "MEDICATION"}
	]
	conditions = [
		item.entity_text
		for item in report.entities
		if str(item.entity_label or "").upper() in {"DISEASE", "SYMPTOM"}
	]
	result = lookup_medicine(
		name,
		other_medicines=other_medicines,
		conditions=conditions,
	)
	entity.entity_text = name
	entity.explanation = ""
	medicine_results = safe_json_list(report.medicine_results)
	result["entity_id"] = entity.id
	for index, saved_result in enumerate(medicine_results):
		if saved_result.get("entity_id") == entity.id:
			medicine_results[index] = result
			break
	else:
		medicine_results.append(result)
	report.medicine_results = json.dumps(medicine_results, ensure_ascii=False)
	db.session.commit()
	log_audit("medicine_lookup", "medical_report", report.id)
	return jsonify(result)


@report_bp.route("/api/reports/<int:report_id>", methods=["PATCH", "DELETE"])
@login_required
def manage_report(report_id):
	report = MedicalReport.query.filter_by(id=report_id, owner_id=get_current_user().id).first_or_404()
	if request.method == "DELETE":
		try:
			delete_report_file(report.file_path, current_app.config["UPLOAD_FOLDER"])
		except OSError as error:
			current_app.logger.error("Report deletion could not remove its file: %s", error)
			return jsonify({"error": "Unable to securely delete the report file."}), 500
		db.session.delete(report)
		db.session.commit()
		log_audit("report_delete", "medical_report", report_id)
		return jsonify({"deleted": True})

	payload = request.get_json(silent=True) or {}
	title = str(payload.get("title", "")).strip()
	if not title:
		return jsonify({"error": "Report title is required."}), 400
	report.title = title[:255]
	db.session.commit()
	log_audit("report_rename", "medical_report", report_id)
	return jsonify(report.to_dict())


@report_bp.route("/api/reports/<int:report_id>/export/<string:file_format>", methods=["GET"])
@login_required
def export_report(report_id, file_format):
	report = MedicalReport.query.filter_by(id=report_id, owner_id=get_current_user().id).first_or_404()
	patient = report.patient
	report_data = report.to_dict()
	data = {
		"patient": {
			"name": patient.full_name if patient else None,
			"age": patient.age if patient else None,
			"gender": patient.gender if patient else None,
			"blood_group": patient.blood_group if patient else None,
		},
		"title": report.title or report.filename or "Medical report",
		"created_at": report.created_at.isoformat() if report.created_at else None,
		"original_text": report.extracted_text,
		"simplified_summary": report_data["simplified_text"],
		"entities": [
			{
				"text": entity.get("entity_text", ""),
				"label": entity.get("entity_label", ""),
				"explanation": entity.get("explanation", ""),
			}
			for entity in report_data["entities"]
		],
		"recommendations": [recommendation.to_dict() for recommendation in report.recommendations],
		"lab_results": report_data["lab_results"],
		"questions_to_ask_doctor": report_data["questions_to_ask_doctor"],
		"disclaimer": DISCLAIMER,
	}
	if file_format == "json":
		return jsonify(data)
	if file_format == "txt":
		lab_lines = []
		if data.get("lab_results"):
			lab_lines = [
				"",
				"Laboratory results:",
				*[
					f"- {lab.get('test_name')}: {lab.get('value')} {lab.get('unit') or ''} (Status: {lab.get('status')})"
					for lab in data["lab_results"]
				],
			]
		questions_lines = []
		if data.get("questions_to_ask_doctor"):
			questions_lines = [
				"",
				"Questions to ask your doctor:",
				*[f"- {q}" for q in data["questions_to_ask_doctor"]],
			]
		content = "\n".join([
			data["title"],
			f"Patient: {data['patient']['name'] or 'Unassigned'}",
			f"Date: {data['created_at'] or 'Unavailable'}",
			"",
			"Simplified summary:",
			data["simplified_summary"] or "No summary available.",
			*lab_lines,
			"",
			"Medical terms:",
			*[
				f"- {item.get('text') or item.get('name')} ({item.get('label') or 'TERM'}): {item.get('explanation') or item.get('general_use') or item.get('use') or 'Prescription medication noted in report.'}"
				for item in data["entities"]
			],
			*questions_lines,
			"",
			data["disclaimer"],
		])
		return send_file(io.BytesIO(content.encode("utf-8")), as_attachment=True, download_name="medical-report.txt", mimetype="text/plain")
	if file_format in {"pdf", "docx"}:
		if file_format == "pdf":
			output = build_simplified_report_pdf(data)
			return send_file(
				output,
				as_attachment=True,
				download_name="medical-report.pdf",
				mimetype="application/pdf",
			)
		from docx import Document
		document = Document()
		document.add_heading(data["title"], level=1)
		document.add_paragraph(f"Patient: {data['patient']['name'] or 'Unassigned'}  |  Date: {data['created_at'] or 'Unavailable'}")
		document.add_heading("Simplified summary", level=2)
		document.add_paragraph(data["simplified_summary"] or "No summary available.")
		if data.get("lab_results"):
			document.add_heading("Laboratory results", level=2)
			for lab in data["lab_results"]:
				document.add_paragraph(f"• {lab.get('test_name')}: {lab.get('value')} {lab.get('unit') or ''} (Status: {lab.get('status')})")
		if data.get("entities"):
			document.add_heading("Medical terms", level=2)
			for item in data["entities"]:
				exp = item.get('explanation') or item.get('general_use') or item.get('use') or 'Prescription medication noted in report.'
				document.add_paragraph(f"• {item.get('text') or item.get('name')}: {exp}")
		if data.get("questions_to_ask_doctor"):
			document.add_heading("Questions to ask your doctor", level=2)
			for q in data["questions_to_ask_doctor"]:
				document.add_paragraph(f"• {q}")
		document.add_paragraph(data["disclaimer"])
		output = io.BytesIO()
		document.save(output)
		output.seek(0)
		return send_file(output, as_attachment=True, download_name="medical-report.docx", mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
	return jsonify({"error": "Supported export formats are pdf, docx, json, and txt."}), 400


@report_bp.route("/api/reports/<int:report_id>/translate", methods=["POST"])
@login_required
def translate_report(report_id):
	report = MedicalReport.query.filter_by(id=report_id, owner_id=get_current_user().id).first_or_404()
	payload = request.get_json(silent=True) or {}
	language = str(payload.get("language", "")).strip()
	if language not in {"en", "ta", "hi"}:
		return jsonify({"error": "Supported languages are en, ta, and hi."}), 400

	report_data = report.to_dict()
	output = {
		"summary": report_data["simplified_text"] or "",
		"entities": [entity.get("explanation") or "" for entity in report_data["entities"]],
		"lab_meanings": [result.get("meaning") or "" for result in report_data["lab_results"]],
		"questions": report_data["questions_to_ask_doctor"],
		"disclaimer": DISCLAIMER,
	}
	try:
		return jsonify({"language": language, "output": translate_report_output(output, language)})
	except TranslationNotConfiguredError as error:
		return jsonify({"error": str(error)}), 503
	except TranslationProviderError as error:
		current_app.logger.warning("Report translation failed: %s", error)
		return jsonify({"error": str(error)}), 502


@report_bp.route("/reports/<int:report_id>/summary", methods=["GET"])
@report_bp.route("/api/reports/<int:report_id>/summary", methods=["GET"])
def public_report_summary(report_id):
	"""Public, scannable, mobile-responsive appointment summary page for QR code visits."""
	report = MedicalReport.query.filter_by(id=report_id).first_or_404()
	patient = report.patient
	report_data = report.to_dict()

	if request.path.startswith("/api/") or "json" in request.headers.get("Accept", ""):
		return jsonify({
			"patient": patient.to_dict() if patient else None,
			"report": report_data,
		})

	return render_template(
		"summary_qr.html",
		report=report,
		patient=patient,
		report_data=report_data,
		disclaimer=DISCLAIMER,
	)


@report_bp.route("/reports/<int:report_id>/download", methods=["GET"])
@login_required
def download_report(report_id):
	report = MedicalReport.query.filter_by(id=report_id, owner_id=get_current_user().id).first_or_404()
	if report.file_path and os.path.exists(report.file_path):
		return send_file(report.file_path, as_attachment=True, download_name=report.filename)
	return abort(404)


@report_bp.route("/reports/<int:report_id>/delete", methods=["POST"])
@login_required
def delete_report(report_id):
	report = MedicalReport.query.filter_by(id=report_id, owner_id=get_current_user().id).first_or_404()
	try:
		delete_report_file(report.file_path, current_app.config["UPLOAD_FOLDER"])
	except OSError as error:
		current_app.logger.error("Report deletion could not remove its file: %s", error)
		return jsonify({"error": "Unable to securely delete the report file."}), 500
	db.session.delete(report)
	db.session.commit()
	log_audit("report_delete", "medical_report", report_id)
	return redirect(url_for("report.list_reports"))


@report_bp.route("/api/admin/audit-logs", methods=["GET"])
@login_required
@require_role("admin")
def audit_logs():
	from database.models import AuditLog
	logs = AuditLog.query.order_by(AuditLog.created_at.desc()).limit(200).all()
	return jsonify([
		{
			"action": log.action,
			"resource_type": log.resource_type,
			"resource_id": log.resource_id,
			"user_id": log.user_id,
			"created_at": log.created_at.isoformat() if log.created_at else None,
		}
		for log in logs
	])


@report_bp.route("/api/admin/backup", methods=["GET"])
@login_required
@require_role("admin")
def backup_database():
	db_path = os.path.join(current_app.config["BASE_DIR"], "database", "medreport.db")
	if not os.path.exists(db_path):
		return jsonify({"error": "Database backup is unavailable."}), 404
	return send_file(db_path, as_attachment=True, download_name="medreport-backup.db", mimetype="application/octet-stream")


@report_bp.route("/api/admin/restore", methods=["POST"])
@login_required
@require_role("admin")
def restore_database():
	backup = request.files.get("backup")
	if not backup or not backup.filename.lower().endswith(".db"):
		return jsonify({"error": "Upload a valid SQLite .db backup."}), 400
	db_path = os.path.join(current_app.config["BASE_DIR"], "database", "medreport.db")
	backup.save(db_path)
	return jsonify({"restored": True})
